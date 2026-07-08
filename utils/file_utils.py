"""Utility helpers for file hashing and text extraction."""

import hashlib
import re
from pathlib import Path

try:
    import fitz
except ImportError:
    fitz = None

try:
    from docx import Document as DocxDoc
except ImportError:
    DocxDoc = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pandas as pd
except ImportError:
    pd = None


def fmt_size(n: int) -> str:
    """Format file size in human-readable form."""
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


def byte_hash(path: Path) -> str:
    """Compute the SHA256 hash of a file's raw bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def extract_text(path: Path) -> str:
    """Extract text from a supported document type."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf" and fitz is not None:
            doc = fitz.open(str(path))
            return " ".join(p.get_text() for p in doc)
        if ext == ".docx" and DocxDoc is not None:
            return " ".join(p.text for p in DocxDoc(str(path)).paragraphs)
        if ext in (".xlsx", ".xls") and openpyxl is not None:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.extend(str(c) for c in row if c)
            return " ".join(out)
        if ext == ".csv" and pd is not None:
            return pd.read_csv(path, dtype=str, nrows=500).to_string()
        if ext in (".md", ".txt", ".csv"):
            return path.read_text(errors="ignore")
    except Exception:
        pass
    return ""


def content_hash(path: Path) -> str:
    """Return a normalized hash for a document's content."""
    text = extract_text(path)
    if not text.strip():
        return byte_hash(path)
    return hashlib.sha256(re.sub(r"\s+", " ", text).strip().lower().encode()).hexdigest()
